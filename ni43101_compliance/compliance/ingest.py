"""Read a report (PDF or DOCX) into a list of Page objects.

PDFs keep true page numbers. DOCX has no fixed pages, so we synthesize
page-sized blocks (~3000 chars) and number them so citations still work.
"""

from __future__ import annotations

import os
from typing import List

from .models import Page

_DOCX_BLOCK_CHARS = 3000


def load_report(path: str) -> List[Page]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext in (".docx", ".doc"):
        return _load_docx(path)
    raise ValueError(f"Unsupported report format: {ext}. Use PDF or DOCX.")


def _load_pdf(path: str) -> List[Page]:
    import fitz  # PyMuPDF

    pages: List[Page] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            pages.append(Page(number=i, text=text))
    return pages


def _load_docx(path: str) -> List[Page]:
    import docx

    document = docx.Document(path)
    parts: List[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full = "\n".join(parts)
    pages: List[Page] = []
    for i in range(0, max(len(full), 1), _DOCX_BLOCK_CHARS):
        block = full[i : i + _DOCX_BLOCK_CHARS]
        pages.append(Page(number=len(pages) + 1, text=block))
    return pages

"""Turn a company compliance checklist (Word or PDF) into atomic Criterion objects.

Checklists vary wildly (prose, bullets, tables), so the default path uses the
LLM to extract discrete, checkable requirements. A non-LLM fallback treats each
bullet / table row as one criterion.
"""

from __future__ import annotations

import os
import re
from typing import List

from .llm import OllamaClient
from .models import Criterion
from .structure import map_criterion_to_item

_EXTRACT_SYSTEM = (
    "You extract discrete compliance requirements from a checklist for NI 43-101 "
    "technical reports. Return ONLY valid JSON of the form "
    '{"criteria": [{"text": "...", "category": "..."}]}. '
    "Each criterion must be a single, self-contained, checkable requirement stated "
    "as an expectation of the report (e.g. 'The report discloses the effective date "
    "of the mineral resource estimate'). Split compound requirements into separate "
    "items. Do not invent requirements; only rephrase what is present. 'category' is "
    "an optional short grouping label if the checklist provides one, else empty."
)


def load_checklist_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        import fitz

        parts = []
        with fitz.open(path) as doc:
            for page in doc:
                parts.append(page.get_text("text") or "")
        return "\n".join(parts)
    if ext in (".docx", ".doc"):
        import docx

        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported checklist format: {ext}. Use PDF or DOCX.")


def extract_criteria(
    raw_text: str,
    llm: OllamaClient,
    use_llm: bool = True,
) -> List[Criterion]:
    if use_llm:
        criteria = _extract_with_llm(raw_text, llm)
    else:
        criteria = _extract_heuristic(raw_text)

    # Attach a best-effort NI 43-101 Item mapping to each.
    for crit in criteria:
        mapped = map_criterion_to_item(crit.text)
        if mapped:
            crit.item_id, crit.item_title = mapped
    return criteria


def _extract_with_llm(raw_text: str, llm: OllamaClient) -> List[Criterion]:
    criteria: List[Criterion] = []
    seen = set()
    # Chunk the checklist so each request stays small (local context is limited).
    for block in _split_blocks(raw_text, max_chars=4000):
        try:
            result = llm.chat_json(_EXTRACT_SYSTEM, f"Checklist text:\n\n{block}")
        except Exception:
            continue
        for item in result.get("criteria", []):
            text = (item.get("text") or "").strip()
            if not text or text.lower() in seen:
                continue
            seen.add(text.lower())
            criteria.append(
                Criterion(
                    id=f"C{len(criteria) + 1:03d}",
                    text=text,
                    category=(item.get("category") or "").strip(),
                )
            )
    if not criteria:
        # LLM produced nothing usable; fall back so the run still works.
        return _extract_heuristic(raw_text)
    return criteria


def _extract_heuristic(raw_text: str) -> List[Criterion]:
    criteria: List[Criterion] = []
    for line in raw_text.splitlines():
        cleaned = re.sub(r"^\s*(?:[-*•●]|\d+[.)]|[a-z][.)])\s*", "", line).strip()
        # Keep lines that look like requirements, not headers or noise.
        if len(cleaned) < 15 or len(cleaned) > 400:
            continue
        criteria.append(Criterion(id=f"C{len(criteria) + 1:03d}", text=cleaned))
    return criteria


def _split_blocks(text: str, max_chars: int) -> List[str]:
    lines = text.splitlines()
    blocks, current, size = [], [], 0
    for line in lines:
        if size + len(line) > max_chars and current:
            blocks.append("\n".join(current))
            current, size = [], 0
        current.append(line)
        size += len(line) + 1
    if current:
        blocks.append("\n".join(current))
    return blocks
